"""Import en masse de liasses fiscales dans le Scoring (API partenaires PVC).

Chaque ligne du fichier Excel crée un dossier, y dépose la liasse et lance son analyse.
Les analyses sont ensuite traitées une par une par le serveur.

Préparation
    python tools/import_liasses.py --modele C:\\import
        -> crée C:\\import\\dossiers.xlsx (une ligne par liasse à remplir)

Import
    python tools/import_liasses.py C:\\import
    python tools/import_liasses.py C:\\import --essai      (vérifie le fichier sans rien envoyer)

Mise à jour de dossiers déjà importés (sans relancer leur analyse)
    python tools/import_liasses.py C:\\import --maj
        -> écrit maj_serveur.sh : commande à coller sur le serveur une fois les analyses terminées

La clé API (PVC_API_KEY du serveur) est lue dans la variable d'environnement WFB_PVC_API_KEY,
sinon demandée au lancement (saisie masquée). Un compte rendu est écrit dans resultats.csv.
"""
from __future__ import annotations

import argparse
import csv
import getpass
import json
import os
import sys
from datetime import datetime
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook

DEFAULT_URL = "http://10.0.30.10:8102/api/v1/partners/pvc/dossiers"
SHEET_COLUMNS = [
    "fichier", "nomClient", "noDemande", "noTiers", "montantTotal", "activite", "fournisseur",
    "duree", "natureBien", "etatBien", "montantHT", "noContrat", "rc", "villeRC",
    "analyste", "dateReception", "urgent",
]
HELP_ROW = [
    "nom du PDF (obligatoire)",
    "raison sociale (obligatoire)",
    "référence unique (vide = générée)",
    "facultatif",
    "facultatif (MAD TTC)",
    "secteur (facultatif)",
    "facultatif",
    "mois (facultatif)",
    "description du bien",
    "neuf / occasion",
    "MAD HT (facultatif)",
    "réf. proforma / contrat",
    "n° RC",
    "ville du RC",
    "nom de l'analyste",
    "JJ/MM/AAAA",
    "oui / non",
]
# Champs du dossier Scoring mis à jour par --maj (jamais le nom ni l'ICE lus dans la liasse).
UPDATED_FIELDS = (
    "sector", "amount", "duration", "valeurBien", "valeurHt", "valeurTtc", "fournisseur",
    "proformaReference", "natureBien", "etat", "rc", "analyst", "date", "urgency",
)

# Exécuté dans le conteneur de l'API : met à jour les dossiers sans relancer d'analyse.
SERVER_SCRIPT = '''import json
from app.schemas.pvc import PvcDossierPv
from app.services import dossier_store
from app.services.pvc_partner import pvc_to_record

FIELDS = __FIELDS__
DATA = json.loads(__DATA__)
for item in DATA:
    existing = dossier_store.get_by_id(item["noDemande"])
    if existing is None:
        print(item["noDemande"], "-> introuvable")
        continue
    if existing.analyseStatus in ("queued", "processing"):
        print(item["noDemande"], "-> analyse en cours : relancez plus tard")
        continue
    built = pvc_to_record(PvcDossierPv.model_validate(item), existing.files, existing=existing)
    update = {field: getattr(built, field) for field in FIELDS}
    if existing.rc:
        update.pop("rc")  # RC déjà connu : conservé
    update["pvc"] = {**(existing.pvc or {}), **(built.pvc or {})}
    dossier_store.prepend(existing.model_copy(update=update))
    print(item["noDemande"], "-> mis à jour :", existing.name, "|", update.get("sector") or "-", "|", update["duration"], "mois")
'''


def write_template(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / "dossiers.xlsx"
    if target.exists():
        sys.exit(f"{target} existe déjà : supprimez-le ou complétez-le.")
    book = Workbook()
    sheet = book.active
    sheet.title = "dossiers"
    sheet.append(SHEET_COLUMNS)
    sheet.append(HELP_ROW)
    sheet.append([
        "BILAN 2025 EXEMPLE.pdf", "STE EXEMPLE SARL", "", "", 250000, "Transport", "FOURNISSEUR EXEMPLE",
        48, "Camion porteur 19 t", "neuf", 208333, "PF-2026-001", "12345", "Casablanca",
        "Analyste Wafabail", "01/10/2026", "non",
    ])
    book.save(target)
    print(f"Modèle créé : {target}")
    print("Remplissez une ligne par liasse (supprimez les lignes d'aide et d'exemple), puis lancez l'import.")


def read_rows(sheet_path: Path) -> list[dict[str, str]]:
    sheet = load_workbook(sheet_path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(cell or "").strip() for cell in rows[0]]
    missing = [col for col in ("fichier", "nomClient") if col not in header]
    if missing:
        sys.exit(f"Colonnes manquantes dans {sheet_path.name} : {', '.join(missing)}")
    result = []
    for values in rows[1:]:
        row = {header[i]: _cell(v) for i, v in enumerate(values) if i < len(header)}
        if not row.get("fichier") and not row.get("nomClient"):
            continue  # ligne vide
        if row.get("fichier") == HELP_ROW[0]:
            continue  # ligne d'aide du modèle
        result.append(row)
    return result


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _amount(value: str, column: str) -> float:
    try:
        return float(value.replace(" ", "").replace("\u202f", "").replace(",", "."))
    except ValueError:
        raise ValueError(f"{column} invalide : {value}") from None


def build_payload(row: dict[str, str], no_demande: str) -> dict:
    """Ligne Excel -> payload DossierPv (mêmes noms de champs que pv-manager)."""
    payload: dict = {"noDemande": no_demande, "nomClient": row["nomClient"]}
    for column, key in (
        ("noTiers", "noTiers"), ("activite", "activite"), ("fournisseur", "fournisseur"),
        ("rc", "rcAnalytique"), ("villeRC", "villeRC"), ("analyste", "commercialNom"),
        ("dateReception", "dateReceptionEbail"),
    ):
        if row.get(column):
            payload[key] = row[column]
    if row.get("montantTotal"):
        payload["montantTotal"] = _amount(row["montantTotal"], "montantTotal")
    if row.get("urgent"):
        payload["enRAC"] = row["urgent"].strip().lower() in {"oui", "o", "yes", "1", "true", "vrai"}

    # Financement : une simulation avec un bien (durée, nature, état, montant HT, référence).
    bien: dict = {}
    if row.get("natureBien"):
        bien["description"] = row["natureBien"]
    if row.get("etatBien"):
        bien["neufOccasion"] = row["etatBien"]
    if row.get("montantHT"):
        bien["montantHT"] = _amount(row["montantHT"], "montantHT")
    simulation: dict = {}
    if row.get("duree"):
        simulation["duree"] = int(round(_amount(row["duree"], "duree")))
    if row.get("noContrat"):
        simulation["noContrat"] = row["noContrat"]
    if bien:
        simulation["biens"] = [bien]
    if simulation:
        payload["simulations"] = [simulation]
    return payload


def write_server_update(folder: Path, payloads: list[dict]) -> Path:
    """Commande à coller sur le serveur : met à jour les dossiers sans relancer d'analyse."""
    code = SERVER_SCRIPT.replace("__FIELDS__", repr(UPDATED_FIELDS)).replace(
        "__DATA__", repr(json.dumps(payloads, ensure_ascii=False))
    )
    target = folder / "maj_serveur.sh"
    target.write_text(
        "# Mise à jour des dossiers importés (sans relancer l'analyse) : à coller sur le serveur\n"
        "sudo docker exec -i scoring-app-api python - <<'EOF'\n" + code + "EOF\n",
        encoding="utf-8",
        newline="\n",
    )
    return target


def main() -> None:
    # Console Windows : ne jamais planter sur un caractère non affichable.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description="Import en masse de liasses fiscales dans le Scoring.")
    parser.add_argument("dossier", type=Path, help="dossier contenant les PDF et dossiers.xlsx")
    parser.add_argument("--modele", action="store_true", help="crée le fichier dossiers.xlsx à remplir")
    parser.add_argument("--fichier", default="dossiers.xlsx", help="nom du fichier Excel (défaut : dossiers.xlsx)")
    parser.add_argument("--url", default=DEFAULT_URL, help=f"adresse de l'API (défaut : {DEFAULT_URL})")
    parser.add_argument("--essai", action="store_true", help="vérifie le fichier et les PDF sans rien envoyer")
    parser.add_argument("--maj", action="store_true", help="prépare la mise à jour serveur des dossiers déjà importés")
    args = parser.parse_args()

    folder: Path = args.dossier
    if args.modele:
        write_template(folder)
        return

    sheet_path = folder / args.fichier
    if not sheet_path.exists():
        sys.exit(f"{sheet_path} introuvable. Créez-le avec : python tools/import_liasses.py --modele {folder}")
    rows = read_rows(sheet_path)
    if not rows:
        sys.exit("Aucune ligne à importer.")

    # Vérification complète avant tout envoi.
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    problems: list[str] = []
    seen: set[str] = set()
    plan = []
    for index, row in enumerate(rows, start=1):
        pdf = folder / row.get("fichier", "")
        no_demande = row.get("noDemande") or f"IMPORT-{stamp}-{index:03d}"
        if args.maj:
            if not row.get("noDemande"):
                problems.append(f"ligne {index} : noDemande obligatoire pour une mise à jour")
        elif not row.get("fichier") or not pdf.is_file():
            problems.append(f"ligne {index} : PDF introuvable ({row.get('fichier') or 'vide'})")
        elif pdf.read_bytes()[:4] != b"%PDF":
            problems.append(f"ligne {index} : {pdf.name} n'est pas un PDF")
        if not row.get("nomClient"):
            problems.append(f"ligne {index} : nomClient vide")
        if no_demande in seen:
            problems.append(f"ligne {index} : noDemande en double ({no_demande})")
        seen.add(no_demande)
        try:
            payload = build_payload(row, no_demande) if row.get("nomClient") else None
        except ValueError as exc:
            problems.append(f"ligne {index} : {exc}")
            payload = None
        plan.append((index, pdf, no_demande, payload))

    if problems:
        print("Fichier à corriger :")
        for problem in problems:
            print(f"  - {problem}")
        sys.exit(1)

    if args.maj:
        target = write_server_update(folder, [payload for _, _, _, payload in plan])
        print(f"{len(plan)} dossier(s) à mettre à jour. Commande serveur écrite dans : {target}")
        print("Ouvrez ce fichier, copiez tout son contenu et collez-le dans le terminal du serveur.")
        return

    print(f"{len(plan)} liasse(s) prêtes à l'envoi.")
    if args.essai:
        for index, pdf, no_demande, payload in plan:
            print(f"  {index:>3}. {no_demande}  {payload['nomClient']}  <-  {pdf.name}")
        print("Essai terminé : rien n'a été envoyé.")
        return

    api_key = os.getenv("WFB_PVC_API_KEY") or getpass.getpass("Clé API (PVC_API_KEY, saisie masquée) : ")
    if not api_key.strip():
        sys.exit("Clé API vide.")

    report_path = folder / "resultats.csv"
    ok = 0
    with httpx.Client(timeout=httpx.Timeout(300.0, connect=15.0)) as client, report_path.open(
        "w", newline="", encoding="utf-8-sig"
    ) as report_file:
        report = csv.writer(report_file, delimiter=";")
        report.writerow(["ligne", "noDemande", "nomClient", "fichier", "resultat", "dossier_scoring", "message"])
        for index, pdf, no_demande, payload in plan:
            label = f"{index:>3}. {no_demande}  {payload['nomClient']}"
            try:
                with pdf.open("rb") as handle:
                    response = client.post(
                        args.url,
                        headers={"X-API-Key": api_key.strip()},
                        data={"data": json.dumps(payload, ensure_ascii=False)},
                        files={"liasse": (pdf.name, handle, "application/pdf")},
                    )
                body = response.json() if "json" in response.headers.get("content-type", "") else {}
                if response.status_code in (200, 201, 202):
                    ok += 1
                    print(f"{label}  ->  créé ({body.get('wfb_id', '')}), analyse en file d'attente")
                    report.writerow([index, no_demande, payload["nomClient"], pdf.name, "créé", body.get("wfb_id", ""), body.get("message", "")])
                else:
                    detail = body.get("detail") or response.text[:200]
                    print(f"{label}  ->  ERREUR {response.status_code} : {detail}")
                    report.writerow([index, no_demande, payload["nomClient"], pdf.name, f"erreur {response.status_code}", "", detail])
                    if response.status_code == 401:
                        sys.exit("Clé API refusée : import arrêté.")
            except httpx.HTTPError as exc:
                print(f"{label}  ->  ERREUR réseau : {exc}")
                report.writerow([index, no_demande, payload["nomClient"], pdf.name, "erreur réseau", "", str(exc)])

    print(f"\n{ok}/{len(plan)} dossier(s) créé(s). Compte rendu : {report_path}")
    print("Les analyses tournent sur le serveur, une à la fois : suivez-les dans la liste des dossiers du Scoring.")


if __name__ == "__main__":
    main()
