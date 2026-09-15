from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from app.sector.domain import ParsedObservation, ParsedWorkbook
from app.sector.mapping import slug_sector_label
from app.sector.registry import SectorDatasetDefinition

_YEAR_RE = re.compile(r"^(20\d{2}|19\d{2})$")
_YEAR_ANNOTATED_RE = re.compile(
    r"^(20\d{2}|19\d{2})(?:\*+|p|e|\(e\)|\(p\)|prov.*)?$",
    re.I,
)
_ISO_DATE_RE = re.compile(r"^(20\d{2}|19\d{2})-\d{2}-\d{2}")
_QUARTER_RE = re.compile(r"(20\d{2}).{0,3}t\s*([1-4])|t\s*([1-4]).{0,3}(20\d{2})", re.I)
_UNIT_RE = re.compile(r"million", re.I)


class SchemaValidationError(ValueError):
    pass


def parse_hcp_number(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace("\xa0", " ").replace(" ", "")
    text = text.replace(",", ".")
    if not text or text in {"-", "…", "...", "n.d.", "nd"}:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def normalize_hcp_label(text: str | None) -> str:
    return " ".join(str(text or "").replace("’", "'").replace("`", "'").split())


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_period(header) -> tuple[int, int | None] | None:
    if isinstance(header, datetime):
        return header.year, None
    if isinstance(header, date):
        return header.year, None
    if isinstance(header, bool):
        return None
    if isinstance(header, (int, float, Decimal)):
        year = int(header)
        if 1990 <= year <= 2100 and abs(float(header) - year) < 1e-6:
            return year, None
        return None
    raw = _cell_text(header)
    if not raw:
        return None
    compact = raw.replace("\xa0", "").replace(" ", "")
    match = _QUARTER_RE.search(compact)
    if match:
        if match.group(1) and match.group(2):
            return int(match.group(1)), int(match.group(2))
        if match.group(3) and match.group(4):
            return int(match.group(4)), int(match.group(3))
    iso = _ISO_DATE_RE.match(raw)
    if iso:
        return int(iso.group(1)), None
    year_only = _YEAR_RE.match(compact) or _YEAR_ANNOTATED_RE.match(compact)
    if year_only:
        return int(year_only.group(1)), None
    return None


def _parse_period(header) -> tuple[int, int | None] | None:
    return parse_period(header)


def detect_structure(rows: list[list]) -> tuple[int, int, list[tuple[int, int, int | None]]]:
    header_idx = None
    periods: list[tuple[int, int, int | None]] = []
    label_col = 0
    for i, row in enumerate(rows[:20]):
        found: list[tuple[int, int, int | None]] = []
        for j, cell in enumerate(row):
            period = parse_period(cell)
            if period:
                found.append((j, period[0], period[1]))
        if len(found) >= 2:
            header_idx = i
            periods = found
            for j, cell in enumerate(row):
                text = _cell_text(cell).lower()
                if text and parse_period(cell) is None:
                    label_col = j
                    break
            break
    if header_idx is None or not periods:
        raise SchemaValidationError("FAILED_SCHEMA: colonnes de période introuvables")
    return header_idx, label_col, periods


def detect_unit(rows: list[list]) -> str:
    blob = " ".join(_cell_text(cell) for row in rows[:6] for cell in row)
    if _UNIT_RE.search(blob):
        return "M MAD"
    return "MAD"


class HcpWorkbookParser:
    def parse(self, content: bytes, definition: SectorDatasetDefinition) -> ParsedWorkbook:
        if content[:2] != b"PK":
            raise SchemaValidationError("FAILED_SCHEMA: fichier non XLSX (signature ZIP absente)")
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        if not rows:
            raise SchemaValidationError("FAILED_SCHEMA: classeur vide")
        header_idx, label_col, periods = detect_structure(rows)
        unit = definition.unit_hint or detect_unit(rows)
        title = _cell_text(rows[0][0]) if rows and rows[0] else definition.name
        observations: list[ParsedObservation] = []
        seen: set[tuple[str, int, int | None]] = set()
        for row in rows[header_idx + 1 :]:
            label = _cell_text(row[label_col] if label_col < len(row) else "")
            if not label or label.lower() == "total":
                continue
            code = slug_sector_label(label)
            # Harmoniser avec le registre de mapping connu
            from app.sector.mapping import HCP_BRANCHES, normalize_activity

            for known_code, known_label in HCP_BRANCHES.items():
                if normalize_activity(label) == normalize_activity(known_label):
                    code = known_code
                    label = known_label
                    break
            for col, year, quarter in periods:
                if col >= len(row):
                    continue
                number = parse_hcp_number(row[col])
                if number is None:
                    continue
                if year < 1990 or year > 2100:
                    raise SchemaValidationError(f"FAILED_SCHEMA: année implausible {year}")
                key = (code, year, quarter)
                if key in seen:
                    raise SchemaValidationError(f"FAILED_SCHEMA: doublon {key}")
                seen.add(key)
                observations.append(
                    ParsedObservation(
                        sector_code=code,
                        sector_label=label,
                        period_type=definition.frequency,
                        year=year,
                        quarter=quarter,
                        value=number,
                        unit=unit,
                        metric=definition.metric,
                        price_type=definition.price_type,
                        base_year=definition.base_year,
                    )
                )
        if not observations:
            raise SchemaValidationError("FAILED_SCHEMA: aucune observation exploitable")
        return ParsedWorkbook(
            observations=observations,
            unit=unit,
            header_labels=[_cell_text(rows[header_idx][c]) for c, _, _ in periods],
            title=title,
            rows_read=len(observations),
        )
